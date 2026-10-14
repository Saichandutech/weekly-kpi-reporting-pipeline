import sqlite3

import pandas as pd
import pytest

from pipeline import etl, report


@pytest.fixture(scope="module")
def db(tmp_path_factory):
    path = tmp_path_factory.mktemp("wh") / "test.db"
    etl.run(db=path)
    return path


def test_test_orders_excluded(db):
    with sqlite3.connect(db) as con:
        assert con.execute("SELECT COUNT(*) FROM fact_orders WHERE channel = 'TEST'").fetchone()[0] == 0


def test_rerun_is_idempotent(db):
    with sqlite3.connect(db) as con:
        before = con.execute("SELECT COUNT(*) FROM fact_orders").fetchone()[0]
    etl.run(db=db)
    with sqlite3.connect(db) as con:
        assert con.execute("SELECT COUNT(*) FROM fact_orders").fetchone()[0] == before


def test_weeks_start_on_monday(db):
    with sqlite3.connect(db) as con:
        weeks = pd.read_sql("SELECT week_start FROM vw_weekly_kpis", con)["week_start"]
    assert (pd.to_datetime(weeks).dt.dayofweek == 0).all()


def test_net_adds_reconcile_to_subscriber_base(db):
    with sqlite3.connect(db) as con:
        k = pd.read_sql("SELECT * FROM vw_weekly_kpis", con)
    assert k["ending_subscribers"].iloc[-1] == etl.STARTING_SUBSCRIBERS + k["net_adds"].sum()


def test_region_totals_match_company_totals(db):
    with sqlite3.connect(db) as con:
        k = pd.read_sql("SELECT week_start, new_connects FROM vw_weekly_kpis", con).set_index("week_start")
        r = pd.read_sql("SELECT week_start, SUM(new_connects) n FROM vw_weekly_orders_by_region GROUP BY 1", con).set_index("week_start")
    assert (k["new_connects"] == r["n"]).all()


def test_callouts_detect_south_texas_issue(db):
    d = report.load(db=db)
    assert any("South Texas" in c and "disconnects" in c for c in d.callouts)


def test_validation_rejects_unknown_region():
    tables = etl.transform(etl.extract())
    tables["fact_orders"].loc[0, "region_code"] = "R9"
    with pytest.raises(etl.DataQualityError):
        etl.validate(tables)


def test_reports_written(db, tmp_path):
    d = report.load(db=db)
    xlsx = report.write_excel(d, tmp_path, db=db)
    assert xlsx.exists()
    from openpyxl import load_workbook
    assert "Data Quality" in load_workbook(xlsx).sheetnames
    assert report.write_html(d, tmp_path).exists()


def test_bi_extracts_match_views(db, tmp_path):
    paths = report.write_bi_extracts(db=db, out=tmp_path)
    kpis = pd.read_csv([p for p in paths if p.name == "weekly_kpis.csv"][0])
    with sqlite3.connect(db) as con:
        assert len(kpis) == con.execute("SELECT COUNT(*) FROM vw_weekly_kpis").fetchone()[0]
